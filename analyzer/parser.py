import re
import pdfplumber
from docx import Document


def extract_text_from_pdf(path):
    parts, links = [], []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
            # Clickable links (e.g. "LinkedIn" text that points to a URL)
            for link in page.hyperlinks:
                uri = link.get("uri")
                if uri:
                    links.append(uri)
    return "\n".join(parts), links


def extract_text_from_docx(path):
    doc = Document(path)
    parts = []

    # Page headers and footers (many templates put name/email/phone here, sometimes inside tables)
    for section in doc.sections:
        for part in (section.header, section.footer):
            parts += [p.text for p in part.paragraphs]
            for table in part.tables:
                for row in table.rows:
                    parts += [cell.text for cell in row.cells]

    parts += [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)

    links = [rel.target_ref for rel in doc.part.rels.values() if "hyperlink" in rel.reltype]
    return "\n".join(parts), links


def clean_text(text):
    text = text.replace("\xa0", " ")
    text = re.sub(r"\(cid:\d+\)", " ", text)                    # icon-font leftovers
    text = re.sub(r"(?<![\w+#])[#+§ª](?=\s)", " ", text)        # standalone icon symbols (keeps C++ / C#)
    text = re.sub(r"\s*\n\s*[•▪●◦]\s*\n?\s*", "\n• ", text)     # fix broken bullets
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def extract_contact_info(text, links):
    search_space = text + "\n" + "\n".join(links)

    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", search_space)
    linkedin = re.search(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w-]+", search_space, re.I)
    github = re.search(r"(?:https?://)?(?:www\.)?github\.com/[\w-]+", search_space, re.I)

    phone = None
    for candidate in re.findall(r"\+?\d[\d\s\-()]{8,}\d", text):
        digits = re.sub(r"\D", "", candidate)
        if 10 <= len(digits) <= 15:
            phone = candidate.strip()
            break

    return {
        "email": email.group().rstrip(".") if email else None,
        "phone": phone,
        "linkedin": linkedin.group() if linkedin else None,
        "github": github.group() if github else None,
    }

SECTION_HEADINGS = {
    "summary": ["summary", "profile", "professional summary", "professional profile", "about me",
                "about", "objective", "career objective", "personal statement", "career summary"],
    "education": ["education", "academic background", "education and training", "academic qualifications",
                  "qualifications", "education & training"],
    "experience": ["experience", "work experience", "professional experience", "employment history",
                   "employment", "work history", "career history", "internships", "internship experience",
                   "relevant experience", "professional background"],
    "projects": ["projects", "academic projects", "personal projects", "key projects", "selected projects"],
    "skills": ["skills", "technical skills", "core skills", "skills & tools", "technologies", "key skills",
               "core competencies", "competencies", "areas of expertise", "expertise", "skills & abilities",
               "skills and abilities", "tools"],
    "certifications": ["certifications", "certificates", "licenses & certifications", "courses",
                       "training", "licenses", "certifications & courses"],
    "languages": ["languages", "language skills"],
    "achievements": ["achievements", "awards", "honors", "awards & achievements", "honors & awards",
                     "accomplishments"],
    "volunteering": ["volunteering", "volunteer experience", "activities", "extracurricular activities",
                     "leadership & activities"],
    "publications": ["publications", "research", "research experience"],
    "references": ["references"],
}

HEADING_LOOKUP = {h: section for section, names in SECTION_HEADINGS.items() for h in names}


def detect_sections(text):
    sections = {}
    current = "header"
    for line in text.split("\n"):
        key = line.strip().lower().rstrip(":").strip()
        if key in HEADING_LOOKUP and len(key) < 40:
            current = HEADING_LOOKUP[key]
            sections.setdefault(current, "")
            continue
        sections[current] = sections.get(current, "") + line + "\n"
    return {name: content.strip() for name, content in sections.items()}


def get_layout_info(path, ext):
    """Detects tables and images — things that often confuse ATS systems."""
    info = {"tables": 0, "images": 0}
    try:
        if ext == "pdf":
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    info["images"] += len(page.images)
                    info["tables"] += len(page.find_tables())
        else:
            doc = Document(path)
            info["tables"] = len(doc.tables)
            info["images"] = len(doc.inline_shapes)
    except Exception:
        pass  # layout info is a bonus — never break the analysis
    return info


def parse_resume(path):
    ext = path.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        raw, links = extract_text_from_pdf(path)
    elif ext == "docx":
        raw, links = extract_text_from_docx(path)
    else:
        raise ValueError("Unsupported file type. Please upload PDF or DOCX.")

    text = clean_text(raw)
    if len(text) < 50:
        raise ValueError(
            "Couldn't read text from this file. It might be a scanned image — "
            "try a text-based PDF or DOCX."
        )

    return {
        "text": text,
        "links": links,
        "word_count": len(text.split()),
        "contact": extract_contact_info(text, links),
        "sections": detect_sections(text),
        "file_type": ext,
        "icon_artifacts": len(re.findall(r"\(cid:\d+\)", raw)),   # counted BEFORE cleaning
        "layout": get_layout_info(path, ext),
    }